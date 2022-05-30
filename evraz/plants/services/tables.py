from typing import Tuple

import openpyxl


class PlantsExcelParser:
    columns_model_fields = {
        'Дата': 'date',
        'Дата выгрузки': 'unload_date',
        'Статус': 'status',
        'Вес(тн)': 'weight',
        '№ УПД': 'upd_number'
    }

    def __init__(self, filename: str):
        self.workbook = openpyxl.load_workbook(filename=filename, data_only=True)

    def generate_plant_from_row(self, row, first_idx: int, last_idx: int) -> Tuple[dict, list]:
        """Валидация и создания словаря со значениями ЗМК, обьектов ЗМК
        соответствующим полям моделей"""
        plant = {}
        plant_objects = []
        plant_object_column_name = 'объект'

        for cell in row[first_idx:last_idx + 1]:
            if not cell.value:
                continue

            column = self.workbook.active[5][cell.col_idx - 1].value

            if column.lower().startswith(plant_object_column_name):
                plant_objects.append(cell.value)

            model_field = self.columns_model_fields.get(column)
            if not model_field:
                continue

            if model_field.endswith('date'):
                plant[model_field] = str(cell.value)[:10]  # cut time, save date
            else:
                plant[model_field] = cell.value

        if not plant.get('date') or not plant.get('unload_date'):
            return {}, []

        return plant, plant_objects

    def get_plants_and_plants_objects(self) -> Tuple[list, list]:
        """Списки данных с ЗМК и их обьектами"""
        ws = self.workbook.active
        plants = []
        plants_objects = []

        for row in ws.iter_rows(6):
            first_idx, last_idx = None, None
            for cell in row:
                if not cell.value:
                    continue

                column_name = ws[5][cell.col_idx-1].value
                model_field = self.columns_model_fields.get(column_name)

                if model_field == 'date':
                    first_idx = cell.col_idx - 1
                elif model_field == 'unload_date':
                    last_idx = cell.col_idx - 1

                if first_idx and last_idx:
                    plant, plant_objects = self.generate_plant_from_row(row, first_idx, last_idx)
                    plants.append(plant)
                    plants_objects.append(plant_objects)
                    first_idx, last_idx = None, None

        return plants, plants_objects
